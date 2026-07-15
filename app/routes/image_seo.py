"""全站图片 SEO 资产、扫描报告和后台管理接口。"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import mimetypes
import re
import struct
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlparse

from flask import Response, jsonify, request, send_file


IMAGE_SUFFIXES = {'.avif', '.gif', '.jpeg', '.jpg', '.png', '.svg', '.webp'}
ROLES = {'primary', 'detail', 'application', 'diagram', 'news', 'decorative', 'logo', 'qrcode'}
NON_INDEXABLE_ROLES = {'decorative', 'logo', 'qrcode'}
ROLE_LABELS = {
    'primary': '产品主图', 'detail': '产品详情图', 'application': '应用场景图',
    'diagram': '原理／结构图', 'news': '新闻资讯图', 'decorative': '页面装饰图',
    'logo': '品牌标识', 'qrcode': '二维码',
}
ROLE_FROM_LABEL = {label: key for key, label in ROLE_LABELS.items()}
ISSUE_LABELS = {
    'missing_alt': '缺少 Alt', 'generic_alt': '泛化 Alt', 'duplicate_alt': '重复 Alt',
    'keyword_stuffing': '关键词重复', 'missing_owner_page': '无归属页面', 'missing_file': '文件不存在',
    'large_file': '文件过大', 'missing_dimensions': '缺少尺寸', 'duplicate_content': '重复内容',
    'unreferenced_asset': '未引用资产',
}
GENERIC_ALTS = {'', 'image', 'img', '图片', '产品图', '新闻图片', 'news', 'product', 'hero'}
MAX_ALT = 220
MAX_TITLE = 220
MAX_CAPTION = 500
MAX_OWNER_PAGE = 500
MAX_URL = 1000
LARGE_IMAGE_BYTES = 2 * 1024 * 1024

_LOCK = threading.RLock()
_CACHE = {'mtime_ns': -1, 'items': []}
_APP_ROOT = Path(__file__).resolve().parents[2]
_DATA_DIR = _APP_ROOT / 'data'
_ASSET_FILE = _DATA_DIR / 'image_seo_assets.json'
_REPORT_FILE = _DATA_DIR / 'image_seo_scan_report.json'
_AUDIT_FILE = _DATA_DIR / 'image_seo_audit.json'


def configure_image_seo(*, app_root=None, data_dir=None):
    global _APP_ROOT, _DATA_DIR, _ASSET_FILE, _REPORT_FILE, _AUDIT_FILE
    if app_root is not None:
        _APP_ROOT = Path(app_root)
    if data_dir is not None:
        _DATA_DIR = Path(data_dir)
    _ASSET_FILE = _DATA_DIR / 'image_seo_assets.json'
    _REPORT_FILE = _DATA_DIR / 'image_seo_scan_report.json'
    _AUDIT_FILE = _DATA_DIR / 'image_seo_audit.json'
    _CACHE['mtime_ns'] = -2
    _CACHE['items'] = []


def _now_iso():
    return datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')


def _read_json(path: Path, default):
    try:
        data = json.loads(path.read_text(encoding='utf-8'))
        return data
    except Exception:
        return default


def _write_json_atomic(path: Path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f'.{path.name}.{uuid.uuid4().hex}.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    tmp.replace(path)


def _clean_text(value, limit):
    text = re.sub(r'\s+', ' ', str(value or '')).strip()
    return text[:limit]


def _normalize_url(value):
    text = _clean_text(value, MAX_URL)
    if not text or text.startswith(('data:', 'blob:', 'javascript:')):
        return ''
    if text.startswith('//'):
        return f'https:{text}'
    return text


def _normalize_owner_page(value):
    text = _clean_text(value, MAX_OWNER_PAGE)
    if not text:
        return ''
    parsed = urlparse(text)
    if parsed.scheme in {'http', 'https'}:
        text = parsed.path or '/'
    if not text.startswith('/'):
        text = '/' + text.lstrip('./')
    return text.split('#', 1)[0]


def _resolve_owner_page_input(value, current=None):
    """Accept an absolute site URL, path, or the asset URL itself."""
    raw = _clean_text(value, MAX_OWNER_PAGE)
    current = current if isinstance(current, dict) else {}
    normalized = _normalize_owner_page(raw)
    parsed = urlparse(raw)
    image_path = urlparse(str(current.get('url') or '')).path
    if parsed.path and image_path and parsed.path == image_path:
        return _normalize_owner_page(current.get('ownerPage') or (current.get('references') or [''])[0])
    if raw and Path(parsed.path or raw).suffix.lower() in IMAGE_SUFFIXES:
        return _normalize_owner_page(current.get('ownerPage') or (current.get('references') or [''])[0])
    return normalized


def _role_for(url: str, alt: str = ''):
    value = f'{url} {alt}'.lower()
    if any(token in value for token in ('qrcode', 'qr-code', '二维码', 'wechat', '微信公众号')):
        return 'qrcode'
    if any(token in value for token in ('logo', 'favicon')):
        return 'logo'
    if any(token in value for token in ('diagram', 'schematic', '原理', '结构图', '示意图')):
        return 'diagram'
    if '/news/' in value or 'news' in value:
        return 'news'
    return 'detail'


def _asset_id(url: str):
    return 'img_' + hashlib.sha256(url.encode('utf-8')).hexdigest()[:20]


def _normalize_record(raw, *, fallback_url=''):
    raw = raw if isinstance(raw, dict) else {}
    url = _normalize_url(raw.get('url') or fallback_url)
    role = str(raw.get('role') or _role_for(url, raw.get('alt'))).strip().lower()
    if role not in ROLES:
        role = 'detail'
    indexable = bool(raw.get('indexable', True)) and role not in NON_INDEXABLE_ROLES
    return {
        'assetId': _clean_text(raw.get('assetId') or _asset_id(url), 80),
        'url': url,
        'ownerPage': _normalize_owner_page(raw.get('ownerPage')),
        'alt': _clean_text(raw.get('alt'), MAX_ALT),
        'title': _clean_text(raw.get('title'), MAX_TITLE),
        'caption': _clean_text(raw.get('caption'), MAX_CAPTION),
        'role': role,
        'ignored': bool(raw.get('ignored', False)),
        'ignoredAt': _clean_text(raw.get('ignoredAt'), 80),
        'indexable': indexable and not bool(raw.get('ignored', False)),
        'width': max(0, int(raw.get('width') or 0)),
        'height': max(0, int(raw.get('height') or 0)),
        'mimeType': _clean_text(raw.get('mimeType'), 120),
        'fileSize': max(0, int(raw.get('fileSize') or 0)),
        'contentHash': _clean_text(raw.get('contentHash'), 128),
        'updatedAt': _clean_text(raw.get('updatedAt') or _now_iso(), 80),
        'references': sorted({_normalize_owner_page(item) for item in raw.get('references', []) if _normalize_owner_page(item)}),
        'issues': sorted({_clean_text(item, 80) for item in raw.get('issues', []) if _clean_text(item, 80)}),
        'sourceAlt': _clean_text(raw.get('sourceAlt'), MAX_ALT),
    }


def load_image_assets():
    with _LOCK:
        try:
            mtime_ns = _ASSET_FILE.stat().st_mtime_ns
        except OSError:
            mtime_ns = -1
        if _CACHE['mtime_ns'] == mtime_ns:
            return [dict(item) for item in _CACHE['items']]
        payload = _read_json(_ASSET_FILE, {'version': 1, 'items': []})
        rows = payload.get('items', []) if isinstance(payload, dict) else []
        items = [_normalize_record(item) for item in rows if isinstance(item, dict) and item.get('url')]
        _CACHE['mtime_ns'] = mtime_ns
        _CACHE['items'] = items
        return [dict(item) for item in items]


def save_image_assets(items):
    normalized = []
    seen = set()
    for item in items or []:
        record = _normalize_record(item)
        if not record['url'] or record['assetId'] in seen:
            continue
        seen.add(record['assetId'])
        normalized.append(record)
    normalized.sort(key=lambda row: (row.get('ownerPage', ''), row.get('url', '')))
    with _LOCK:
        _write_json_atomic(_ASSET_FILE, {'version': 1, 'updatedAt': _now_iso(), 'items': normalized})
        try:
            _CACHE['mtime_ns'] = _ASSET_FILE.stat().st_mtime_ns
        except OSError:
            _CACHE['mtime_ns'] = -1
        _CACHE['items'] = normalized
    return normalized


def image_asset_map():
    result = {}
    for item in load_image_assets():
        result[item['url']] = item
        parsed = urlparse(item['url'])
        if parsed.path:
            result[parsed.path] = item
    return result


def get_image_asset(url: str, owner_page: str = ''):
    normalized_url = _normalize_url(url)
    path = urlparse(normalized_url).path if normalized_url else ''
    owner = _normalize_owner_page(owner_page)
    candidates = []
    for item in load_image_assets():
        if item['url'] not in {normalized_url, path} and urlparse(item['url']).path != path:
            continue
        candidates.append(item)
    if owner:
        for item in candidates:
            if item.get('ownerPage') == owner:
                return item
    return candidates[0] if candidates else None


def get_indexable_images_for_page(owner_page: str):
    owner = _normalize_owner_page(owner_page)
    return [item for item in load_image_assets() if item.get('ownerPage') == owner and item.get('indexable') and not item.get('ignored')]


def register_uploaded_image(url: str, file_path, *, owner_page: str = '', role: str = 'detail'):
    path = Path(file_path)
    if not path.is_file() or path.suffix.lower() not in IMAGE_SUFFIXES:
        return None
    items = load_image_assets()
    normalized_url = _normalize_url(url)
    by_url = {item['url']: item for item in items}
    record = _normalize_record({
        **by_url.get(normalized_url, {}),
        'url': normalized_url,
        'ownerPage': owner_page,
        'role': role,
        'indexable': role not in NON_INDEXABLE_ROLES,
        **_file_metadata(path),
        'updatedAt': _now_iso(),
        'issues': ['missing_alt'] if role not in NON_INDEXABLE_ROLES else [],
    })
    by_url[normalized_url] = record
    save_image_assets(list(by_url.values()))
    return record


def _html_attrs(tag):
    attrs = {}
    for match in re.finditer(r'''(?is)\b([a-z_:][-a-z0-9_:.]*)\s*=\s*(?:"([^"]*)"|'([^']*)'|([^\s>]+))''', tag or ''):
        attrs[match.group(1).lower()] = html.unescape(next((v for v in match.groups()[1:] if v is not None), ''))
    return attrs


def _public_path_for_file(path: Path):
    try:
        rel = path.resolve().relative_to(_APP_ROOT.resolve()).as_posix()
    except Exception:
        return ''
    return '/' if rel == 'index.html' else '/' + rel


def _resolve_local_file(url: str, page_file: Path | None = None):
    parsed = urlparse(url)
    if parsed.scheme in {'http', 'https'}:
        return None
    path = unquote(parsed.path or '')
    if not path:
        return None
    candidates = []
    if path.startswith('/'):
        candidates.append(_APP_ROOT / path.lstrip('/'))
    elif page_file is not None:
        candidates.append(page_file.parent / path)
    candidates.append(_APP_ROOT / path.lstrip('./'))
    root = _APP_ROOT.resolve()
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
            if str(resolved).startswith(str(root)) and resolved.is_file():
                return resolved
        except Exception:
            continue
    return None


def _image_dimensions(path: Path):
    try:
        data = path.read_bytes()[:262144]
    except Exception:
        return 0, 0
    suffix = path.suffix.lower()
    try:
        if suffix == '.png' and data.startswith(b'\x89PNG'):
            return struct.unpack('>II', data[16:24])
        if suffix == '.gif' and data[:6] in {b'GIF87a', b'GIF89a'}:
            return struct.unpack('<HH', data[6:10])
        if suffix in {'.jpg', '.jpeg'} and data.startswith(b'\xff\xd8'):
            pos = 2
            while pos + 9 < len(data):
                if data[pos] != 0xFF:
                    pos += 1
                    continue
                marker = data[pos + 1]
                length = int.from_bytes(data[pos + 2:pos + 4], 'big')
                if marker in range(0xC0, 0xC4) or marker in range(0xC5, 0xC8) or marker in range(0xC9, 0xCC) or marker in range(0xCD, 0xD0):
                    return int.from_bytes(data[pos + 7:pos + 9], 'big'), int.from_bytes(data[pos + 5:pos + 7], 'big')
                pos += max(2, length + 2)
        if suffix == '.webp' and data.startswith(b'RIFF') and data[8:12] == b'WEBP':
            kind = data[12:16]
            if kind == b'VP8 ' and len(data) >= 30:
                # Lossy WebP: the frame header begins after the RIFF/chunk
                # headers. The 0x9d012a start code is followed by two
                # little-endian 14-bit canvas dimensions.
                start_code = data.find(b'\x9d\x01\x2a', 20, min(len(data), 64))
                if start_code >= 0 and start_code + 7 <= len(data):
                    width = int.from_bytes(data[start_code + 3:start_code + 5], 'little') & 0x3FFF
                    height = int.from_bytes(data[start_code + 5:start_code + 7], 'little') & 0x3FFF
                    return width, height
            if kind == b'VP8X' and len(data) >= 30:
                return 1 + int.from_bytes(data[24:27], 'little'), 1 + int.from_bytes(data[27:30], 'little')
            if kind == b'VP8L' and len(data) >= 25:
                bits = int.from_bytes(data[21:25], 'little')
                return (bits & 0x3FFF) + 1, ((bits >> 14) & 0x3FFF) + 1
        if suffix == '.svg':
            text = data.decode('utf-8', errors='ignore')
            w = re.search(r'\bwidth=["\']([0-9.]+)', text)
            h = re.search(r'\bheight=["\']([0-9.]+)', text)
            return int(float(w.group(1))) if w else 0, int(float(h.group(1))) if h else 0
    except Exception:
        return 0, 0
    return 0, 0


def _file_metadata(path: Path | None):
    if path is None:
        return {'width': 0, 'height': 0, 'mimeType': '', 'fileSize': 0, 'contentHash': ''}
    try:
        size = path.stat().st_size
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
    except Exception:
        size, digest = 0, ''
    width, height = _image_dimensions(path)
    mime = mimetypes.guess_type(path.name)[0] or 'application/octet-stream'
    return {'width': width, 'height': height, 'mimeType': mime, 'fileSize': size, 'contentHash': digest}


def _page_type(page: str):
    if '/gassensing/' in page or '/biosensing/' in page or '/customization/' in page:
        return 'product'
    if '/news/' in page:
        return 'news'
    if '/solutions/' in page:
        return 'solution'
    if '/cases/' in page:
        return 'case'
    if page == '/':
        return 'home'
    return 'page'


def _issues_for(record, alt_counts, hash_counts, referenced=True, local_exists=True):
    issues = []
    alt = record.get('alt', '').strip()
    if record.get('role') not in NON_INDEXABLE_ROLES:
        if not alt:
            issues.append('missing_alt')
        elif alt.lower() in GENERIC_ALTS:
            issues.append('generic_alt')
        elif alt_counts.get(alt, 0) > 1:
            issues.append('duplicate_alt')
        words = [word for word in re.split(r'[\s,，。/|]+', alt.lower()) if len(word) >= 2]
        if words and max(words.count(word) for word in set(words)) >= 3:
            issues.append('keyword_stuffing')
    if not record.get('ownerPage'):
        issues.append('missing_owner_page')
    if record.get('url', '').startswith('/') and not local_exists:
        issues.append('missing_file')
    if record.get('fileSize', 0) > LARGE_IMAGE_BYTES:
        issues.append('large_file')
    if referenced and (not record.get('width') or not record.get('height')):
        issues.append('missing_dimensions')
    if record.get('contentHash') and hash_counts.get(record['contentHash'], 0) > 1:
        issues.append('duplicate_content')
    if not referenced:
        issues.append('unreferenced_asset')
    return sorted(set(issues))


def scan_image_assets(*, persist=True):
    existing = {item['url']: item for item in load_image_assets()}
    references = {}
    html_files = [_APP_ROOT / 'index.html'] + sorted((_APP_ROOT / 'pages').rglob('*.html'))
    for page_file in html_files:
        if not page_file.is_file():
            continue
        try:
            body = page_file.read_text(encoding='utf-8')
        except Exception:
            continue
        page = _public_path_for_file(page_file)
        for match in re.finditer(r'(?is)<img\b[^>]*>', body):
            attrs = _html_attrs(match.group(0))
            url = _normalize_url(attrs.get('src'))
            if not url or '${' in url:
                continue
            local_file = _resolve_local_file(url, page_file)
            public_url = _public_path_for_file(local_file) if local_file else url
            if not public_url:
                public_url = url
            row = references.setdefault(public_url, {'pages': set(), 'alts': [], 'file': local_file})
            row['pages'].add(page)
            row['alts'].append(_clean_text(attrs.get('alt'), MAX_ALT))

    asset_files = []
    for dirname in ('assets', 'cdn_assets', 'pages'):
        root = _APP_ROOT / dirname
        if root.exists():
            asset_files.extend(path for path in root.rglob('*') if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES)
    for path in asset_files:
        references.setdefault(_public_path_for_file(path), {'pages': set(), 'alts': [], 'file': path})

    records = []
    for url, ref in references.items():
        old = existing.get(url, {})
        source_alt = next((value for value in ref['alts'] if value), '')
        pages = sorted(ref['pages'])
        role = old.get('role') or _role_for(url, source_alt)
        record = _normalize_record({
            **old,
            'url': url,
            'ownerPage': old.get('ownerPage') or (pages[0] if pages else ''),
            'alt': old.get('alt') if 'alt' in old else source_alt,
            'role': role,
            'indexable': old.get('indexable', role not in NON_INDEXABLE_ROLES),
            'references': pages,
            'sourceAlt': source_alt,
            **_file_metadata(ref['file']),
            'updatedAt': old.get('updatedAt') or _now_iso(),
        })
        record['pageType'] = _page_type(record.get('ownerPage') or (pages[0] if pages else ''))
        records.append(record)

    alt_counts = {}
    hash_counts = {}
    for record in records:
        if record.get('alt'):
            alt_counts[record['alt']] = alt_counts.get(record['alt'], 0) + 1
        if record.get('contentHash'):
            hash_counts[record['contentHash']] = hash_counts.get(record['contentHash'], 0) + 1
    for record in records:
        record['issues'] = _issues_for(
            record,
            alt_counts,
            hash_counts,
            referenced=bool(record.get('references')),
            local_exists=bool(_resolve_local_file(record['url'])),
        )

    summary = {'total': len(records), 'indexable': 0, 'referenced': 0, 'withIssues': 0, 'ignored': 0, 'issues': {}}
    for record in records:
        summary['indexable'] += int(bool(record.get('indexable')))
        summary['referenced'] += int(bool(record.get('references')))
        summary['ignored'] += int(bool(record.get('ignored')))
        summary['withIssues'] += int(bool(record.get('issues')) and not record.get('ignored'))
        for issue in record.get('issues', []):
            summary['issues'][issue] = summary['issues'].get(issue, 0) + 1
    report = {'version': 1, 'generatedAt': _now_iso(), 'summary': summary, 'items': records}
    if persist:
        save_image_assets(records)
        with _LOCK:
            _write_json_atomic(_REPORT_FILE, report)
    return report


def _append_audit(action, details):
    with _LOCK:
        payload = _read_json(_AUDIT_FILE, {'items': []})
        items = payload.get('items', []) if isinstance(payload, dict) else []
        items.append({'timestamp': _now_iso(), 'action': action, 'details': details})
        _write_json_atomic(_AUDIT_FILE, {'items': items[-1000:]})


def _queue_owner_page_submission(items):
    paths = sorted({item.get('ownerPage') for item in items or [] if item.get('ownerPage')})
    if not paths:
        return
    try:
        from app.routes.product_settings import queue_search_engine_submission
        queue_search_engine_submission(paths)
    except Exception:
        # Submission is optional and must never block metadata persistence.
        return


def _draft_alt(item):
    if item.get('role') in NON_INDEXABLE_ROLES:
        return ''
    owner = Path(item.get('ownerPage') or '').stem.replace('_', ' ').replace('-', ' ').strip()
    filename = Path(urlparse(item.get('url') or '').path).stem.replace('_', ' ').replace('-', ' ').strip()
    base = owner if owner and owner not in {'index', 'about'} else filename
    role_labels = {'primary': '主图', 'detail': '详情图', 'application': '应用场景图', 'diagram': '原理示意图', 'news': '新闻图片'}
    return _clean_text(f'{base} {role_labels.get(item.get("role"), "图片")}', MAX_ALT)


def register_image_seo_routes(app, *, login_required, app_root, data_dir, is_same_origin_request=lambda _request: True):
    configure_image_seo(app_root=app_root, data_dir=data_dir)

    def require_same_origin():
        try:
            return bool(is_same_origin_request(request))
        except Exception:
            return False

    @app.route('/api/admin/image-seo/assets', methods=['GET'])
    @login_required
    def list_image_seo_assets():
        items = load_image_assets()
        if not items and request.args.get('scan_if_empty', '1') != '0':
            items = scan_image_assets(persist=True)['items']
        role = str(request.args.get('role') or '').strip()
        issue = str(request.args.get('issue') or '').strip()
        page_type = str(request.args.get('page_type') or '').strip()
        query = str(request.args.get('q') or '').strip().lower()
        indexable = str(request.args.get('indexable') or '').strip().lower()
        status = str(request.args.get('status') or 'active').strip().lower()
        filtered = []
        for item in items:
            row = dict(item)
            row['pageType'] = _page_type(row.get('ownerPage'))
            if status == 'ignored' and not row.get('ignored'):
                continue
            if status != 'ignored' and row.get('ignored'):
                continue
            if role and row.get('role') != role:
                continue
            if issue and issue not in row.get('issues', []):
                continue
            if page_type and row.get('pageType') != page_type:
                continue
            if indexable in {'true', 'false'} and bool(row.get('indexable')) != (indexable == 'true'):
                continue
            if query and query not in ' '.join(str(row.get(k) or '') for k in ('url', 'ownerPage', 'alt', 'title', 'caption')).lower():
                continue
            filtered.append(row)
        page = max(1, int(request.args.get('page') or 1))
        page_size = min(200, max(10, int(request.args.get('page_size') or 50)))
        start = (page - 1) * page_size
        return jsonify({'success': True, 'total': len(filtered), 'page': page, 'page_size': page_size, 'items': filtered[start:start + page_size]})

    @app.route('/api/admin/image-seo/scan', methods=['POST'])
    @login_required
    def run_image_seo_scan():
        if not require_same_origin():
            return jsonify({'success': False, 'message': '请求来源校验失败'}), 403
        report = scan_image_assets(persist=True)
        _append_audit('scan', {'summary': report['summary']})
        return jsonify({'success': True, 'report': report})

    @app.route('/api/admin/image-seo/report', methods=['GET'])
    @login_required
    def get_image_seo_report():
        report = _read_json(_REPORT_FILE, {})
        if not report:
            report = scan_image_assets(persist=True)
        else:
            # Metadata edits (ignore/restore, Alt, role) update the asset
            # store immediately; keep the dashboard summary fresh even
            # before the next full filesystem scan.
            current_items = load_image_assets()
            summary = report.setdefault('summary', {})
            summary['total'] = len(current_items)
            summary['ignored'] = sum(1 for item in current_items if item.get('ignored'))
            summary['indexable'] = sum(1 for item in current_items if item.get('indexable') and not item.get('ignored'))
            summary['referenced'] = sum(1 for item in current_items if item.get('references'))
            summary['withIssues'] = sum(1 for item in current_items if item.get('issues') and not item.get('ignored'))
        return jsonify({'success': True, 'report': report})

    @app.route('/api/admin/image-seo/report.csv', methods=['GET'])
    @login_required
    def download_image_seo_report_csv():
        items = load_image_assets()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['资产编号', '图片显示', '图片地址', '归属页面', '图片角色', '允许索引', '已忽略', '忽略时间', 'Alt替代文本', '图片标题', '图片说明', '宽度', '高度', '文件类型', '文件大小（字节）', '内容哈希', '问题'])
        for item in items:
            writer.writerow([
                item.get('assetId', ''), item.get('url', ''), item.get('url', ''), item.get('ownerPage', ''), ROLE_LABELS.get(item.get('role'), item.get('role', '')),
                '是' if item.get('indexable') else '否', '是' if item.get('ignored') else '否', item.get('ignoredAt', ''),
                item.get('alt', ''), item.get('title', ''), item.get('caption', ''), item.get('width', 0), item.get('height', 0),
                item.get('mimeType', ''), item.get('fileSize', 0), item.get('contentHash', ''), '；'.join(ISSUE_LABELS.get(issue, issue) for issue in item.get('issues', [])),
            ])
        return Response('\ufeff' + output.getvalue(), mimetype='text/csv', headers={'Content-Disposition': 'attachment; filename="image-seo-report.csv"'})

    @app.route('/api/admin/image-seo/report.xlsx', methods=['GET'])
    @login_required
    def download_image_seo_report_xlsx():
        try:
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image as ExcelImage
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.worksheet.datavalidation import DataValidation
            from PIL import Image as PillowImage
        except Exception:
            return jsonify({'success': False, 'message': '服务器尚未安装 Excel 导出依赖'}), 503

        items = load_image_assets()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '图片SEO'
        headers = ['资产编号', '图片显示', '图片地址', '归属页面', '图片角色', '允许索引', '已忽略', '忽略时间', 'Alt替代文本', '图片标题', '图片说明', '宽度', '高度', '文件类型', '文件大小（字节）', '内容哈希', '问题']
        sheet.append(headers)
        header_fill = PatternFill('solid', fgColor='1B648C')
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:Q{max(1, len(items) + 1)}'
        widths = {'A': 25, 'B': 18, 'C': 55, 'D': 48, 'E': 18, 'F': 12, 'G': 12, 'H': 24, 'I': 42, 'J': 34, 'K': 48, 'L': 10, 'M': 10, 'N': 18, 'O': 18, 'P': 68, 'Q': 42}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        image_streams = []
        for row_no, item in enumerate(items, start=2):
            sheet.append([
                item.get('assetId', ''), '', item.get('url', ''), item.get('ownerPage', ''), ROLE_LABELS.get(item.get('role'), item.get('role', '')),
                '是' if item.get('indexable') else '否', '是' if item.get('ignored') else '否', item.get('ignoredAt', ''),
                item.get('alt', ''), item.get('title', ''), item.get('caption', ''), item.get('width', 0), item.get('height', 0),
                item.get('mimeType', ''), item.get('fileSize', 0), item.get('contentHash', ''), '；'.join(ISSUE_LABELS.get(issue, issue) for issue in item.get('issues', [])),
            ])
            sheet.row_dimensions[row_no].height = 78
            for cell in sheet[row_no]:
                cell.alignment = Alignment(vertical='center', wrap_text=True)
            local_file = _resolve_local_file(item.get('url', ''))
            if not local_file:
                continue
            try:
                with PillowImage.open(local_file) as source:
                    source.thumbnail((140, 92))
                    if source.mode not in {'RGB', 'RGBA'}:
                        source = source.convert('RGBA')
                    stream = io.BytesIO()
                    source.save(stream, format='PNG')
                    stream.seek(0)
                image_streams.append(stream)
                preview = ExcelImage(stream)
                preview.anchor = f'B{row_no}'
                sheet.add_image(preview)
            except Exception:
                sheet.cell(row_no, 2, item.get('url', ''))

        role_validation = DataValidation(type='list', formula1='"' + ','.join(ROLE_LABELS.values()) + '"', allow_blank=False)
        yes_no_validation = DataValidation(type='list', formula1='"是,否"', allow_blank=False)
        sheet.add_data_validation(role_validation)
        sheet.add_data_validation(yes_no_validation)
        if items:
            role_validation.add(f'E2:E{len(items) + 1}')
            yes_no_validation.add(f'F2:G{len(items) + 1}')
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name='image-seo-report.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.route('/api/admin/image-seo/import-csv', methods=['POST'])
    @login_required
    def import_image_seo_csv():
        if not require_same_origin():
            return jsonify({'success': False, 'message': '请求来源校验失败'}), 403
        upload = request.files.get('file')
        if not upload or not upload.filename:
            return jsonify({'success': False, 'message': '请选择 CSV 文件'}), 400
        filename = str(upload.filename).lower()
        if not filename.endswith(('.csv', '.xlsx')):
            return jsonify({'success': False, 'message': '只支持 CSV 或 XLSX 文件'}), 400
        try:
            raw = upload.read()
            if filename.endswith('.xlsx'):
                try:
                    from openpyxl import load_workbook
                except Exception:
                    return jsonify({'success': False, 'message': '服务器尚未安装 Excel 导入依赖'}), 503
                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                sheet = workbook['图片SEO'] if '图片SEO' in workbook.sheetnames else workbook.active
                values = sheet.iter_rows(values_only=True)
                headers = [str(value or '').strip() for value in next(values, [])]
                reader = ({headers[index]: value for index, value in enumerate(row) if index < len(headers)} for row in values)
                fields = set(headers)
            else:
                text = raw.decode('utf-8-sig')
                csv_reader = csv.DictReader(io.StringIO(text))
                if not csv_reader.fieldnames:
                    return jsonify({'success': False, 'message': 'CSV 缺少表头'}), 400
                reader = csv_reader
                fields = {str(field or '').strip() for field in csv_reader.fieldnames}
            if not fields.intersection({'资产编号', 'assetId'}) and not fields.intersection({'图片地址', 'url'}):
                return jsonify({'success': False, 'message': 'CSV 必须包含“资产编号”或“图片地址”列'}), 400
            items = load_image_assets()
            by_id = {item['assetId']: item for item in items}
            by_url = {item['url']: item for item in items}
            changed = []
            errors = []
            for row_no, raw_row in enumerate(reader, start=2):
                row = {str(key or '').strip(): value for key, value in raw_row.items()}
                asset_id = _clean_text(row.get('资产编号') or row.get('assetId'), 80)
                url = _normalize_url(row.get('图片地址') or row.get('url') or row.get('图片显示'))
                current = by_id.get(asset_id) if asset_id else by_url.get(url)
                if not current and url:
                    current = next((item for item in items if urlparse(item['url']).path == urlparse(url).path), None)
                if not current:
                    errors.append(f'第 {row_no} 行找不到对应图片资产')
                    continue
                merged = dict(current)
                role = row.get('图片角色') or row.get('role')
                if role:
                    merged['role'] = ROLE_FROM_LABEL.get(str(role).strip(), str(role).strip())
                for target, aliases in {
                    'ownerPage': ('归属页面', 'ownerPage'), 'alt': ('Alt替代文本', 'alt'),
                    'title': ('图片标题', 'title'), 'caption': ('图片说明', 'caption'),
                }.items():
                    for alias in aliases:
                        if alias in row:
                            merged[target] = row.get(alias, '')
                            break
                if 'ownerPage' in merged:
                    merged['ownerPage'] = _resolve_owner_page_input(merged.get('ownerPage'), current)
                for target, aliases in {'indexable': ('允许索引', 'indexable'), 'ignored': ('已忽略', 'ignored')}.items():
                    for alias in aliases:
                        if alias in row and str(row.get(alias, '')).strip() != '':
                            value = str(row.get(alias, '')).strip().lower() in {'是', 'true', '1', 'yes', 'y', 'on'}
                            merged[target] = value
                            break
                if merged.get('ignored'):
                    merged['indexable'] = False
                merged['updatedAt'] = _now_iso()
                normalized = _normalize_record(merged)
                by_id[current['assetId']] = normalized
                changed.append(current['assetId'])
            if changed:
                saved = save_image_assets(list(by_id.values()))
                _append_audit('import_csv', {'assetIds': changed, 'errors': errors[:50]})
                _queue_owner_page_submission([item for item in saved if item['assetId'] in set(changed)])
            return jsonify({'success': True, 'updated': len(set(changed)), 'errors': errors[:50], 'errorCount': len(errors)})
        except UnicodeDecodeError:
            return jsonify({'success': False, 'message': 'CSV 必须使用 UTF-8 编码'}), 400
        except Exception as exc:
            return jsonify({'success': False, 'message': f'CSV 解析失败：{type(exc).__name__}'}), 400

    def update_records(changes):
        items = load_image_assets()
        by_id = {item['assetId']: item for item in items}
        changed = []
        for raw in changes:
            asset_id = _clean_text(raw.get('assetId'), 80)
            current = by_id.get(asset_id)
            if not current:
                continue
            merged = dict(current)
            for key in ('alt', 'title', 'caption', 'ownerPage', 'role', 'indexable', 'ignored'):
                if key in raw:
                    merged[key] = raw[key]
            if 'ownerPage' in raw:
                merged['ownerPage'] = _resolve_owner_page_input(raw.get('ownerPage'), current)
            if 'ignored' in raw:
                merged['ignoredAt'] = _now_iso() if bool(raw['ignored']) else ''
                if bool(raw['ignored']):
                    merged['indexable'] = False
                elif 'indexable' not in raw:
                    merged['indexable'] = str(merged.get('role') or '') not in NON_INDEXABLE_ROLES
            merged['updatedAt'] = _now_iso()
            normalized = _normalize_record(merged)
            by_id[asset_id] = normalized
            changed.append(asset_id)
        if changed:
            saved = save_image_assets(list(by_id.values()))
            _append_audit('update', {'assetIds': changed})
            _queue_owner_page_submission([item for item in saved if item['assetId'] in changed])
        return changed

    @app.route('/api/admin/image-seo/assets/<asset_id>', methods=['PUT'])
    @login_required
    def update_image_seo_asset(asset_id):
        if not require_same_origin():
            return jsonify({'success': False, 'message': '请求来源校验失败'}), 403
        data = request.get_json(silent=True) or {}
        data['assetId'] = asset_id
        changed = update_records([data])
        if not changed:
            return jsonify({'success': False, 'message': '图片资产不存在'}), 404
        return jsonify({'success': True, 'item': next(item for item in load_image_assets() if item['assetId'] == asset_id)})

    @app.route('/api/admin/image-seo/assets/bulk', methods=['POST'])
    @login_required
    def bulk_update_image_seo_assets():
        if not require_same_origin():
            return jsonify({'success': False, 'message': '请求来源校验失败'}), 403
        data = request.get_json(silent=True) or {}
        asset_ids = [str(item) for item in data.get('assetIds', []) if str(item).strip()][:500]
        patch = data.get('patch') if isinstance(data.get('patch'), dict) else {}
        allowed_patch = {key: patch[key] for key in ('role', 'indexable', 'ownerPage', 'ignored') if key in patch}
        if not asset_ids or not allowed_patch:
            return jsonify({'success': False, 'message': '请选择图片并提供允许的批量字段'}), 400
        changed = update_records([{'assetId': asset_id, **allowed_patch} for asset_id in asset_ids])
        return jsonify({'success': True, 'updated': len(changed), 'assetIds': changed})

    @app.route('/api/admin/image-seo/alt-drafts', methods=['POST'])
    @login_required
    def generate_image_alt_drafts():
        if not require_same_origin():
            return jsonify({'success': False, 'message': '请求来源校验失败'}), 403
        data = request.get_json(silent=True) or {}
        selected = {str(item) for item in data.get('assetIds', [])}
        items = load_image_assets()
        drafts = [{'assetId': item['assetId'], 'alt': _draft_alt(item)} for item in items if item['assetId'] in selected]
        return jsonify({'success': True, 'drafts': drafts, 'persisted': False})


__all__ = [
    'configure_image_seo', 'get_image_asset', 'get_indexable_images_for_page',
    'image_asset_map', 'load_image_assets', 'register_image_seo_routes', 'register_uploaded_image', 'scan_image_assets',
]
